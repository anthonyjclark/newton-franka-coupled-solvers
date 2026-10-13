"""Insert the "what is Newton" introduction slides into the demo deck.

    uvx --from python-pptx python slides/add_intro_slides.py in.pptx slides-<timestamp>-<message>.pptx

Adds four slides after "What you will see": what Newton is, what it is not, how it relates to
MuJoCo Warp / Isaac Lab / Isaac Sim, and the repo examples as a starting point. Content comes from
the Newton overview page, the NVIDIA "Getting started with Newton" course, and the examples
shipped in the installed package (`python -m newton.examples --list`). It also sets the agenda
slide's total time to match the script. Refuses to run on a deck that already has them.
"""

import argparse
from pathlib import Path

from deckkit import Kit, title_of
from pptx import Presentation

AFTER = "What you will see"
FIRST = "What is Newton?"

IS = [
    "A GPU-accelerated, extensible physics engine for robotics and simulation research",
    "Built on NVIDIA Warp; maintained by Disney Research, Google DeepMind and NVIDIA",
    "GPU-first: many worlds stepped in parallel on one GPU",
    "Modular: one Model / State / Control / Contacts API, several swappable solvers "
    "(MuJoCo, VBD, XPBD, Featherstone, Kamino, ...)",
    "Imports URDF, MJCF and USD; some solvers are differentiable",
]
IS_NOT = [
    "A photorealistic renderer: its viewers are for debugging and playback",
    "A robot-learning framework: no RL tasks, training loops or policies (Isaac Lab adds those)",
    "A one-solver-fits-all engine: you pick, or couple, solvers per scene",
    "A controller or task library: the keyframes and IK in this demo are our own code on top",
    "Stable everywhere: coupling, VBD and Kamino are experimental",
]
NEIGHBOURS = {
    "cols": [2.4, 4.6, 4.9],
    "rows": [
        ("Project", "What it is", "Relationship to Newton"),
        ("MuJoCo Warp", "GPU MuJoCo from Google DeepMind", "Integrated as SolverMuJoCo (our arm)"),
        ("Isaac Lab", "NVIDIA's robot-learning framework", "Can use Newton as its physics engine for RL"),
        ("Isaac Sim", "NVIDIA's robot simulator (Omniverse, RTX)", "Rendering and scene tooling; Newton stays standalone"),
    ],
}
EXAMPLES_CMD = [
    "$ uv run python -m newton.examples --list",
    "$ uv run python -m newton.examples basic_pendulum",
    "$ uv run python -m newton.examples mujoco_vbd_coupled_solver",
]
EXAMPLES = [
    "basic_*: shapes, joints, URDF import, viewers",
    "cable_*, cloth_*, mpm_*: deformables and granular materials",
    "ik_franka, ik_h1: inverse kinematics",
    "diffsim_*: differentiable simulation",
    "multiphysics: coupled solvers (this demo's cousins: mujoco_vbd_coupled_solver, franka_cable_ik_pick_place)",
    "NVIDIA's course \"Getting started with Newton\" teaches the same concepts step by step",
]


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck", type=Path)
    ap.add_argument("out", type=Path)
    args = ap.parse_args()
    if args.out.resolve() == args.deck.resolve():
        ap.error("out must differ from deck (keep the original)")

    prs = Presentation(args.deck)
    titles = [title_of(s) for s in prs.slides]
    if FIRST in titles:
        ap.error(f"deck already has the {FIRST!r} slide")
    if AFTER not in titles:
        ap.error(f"no slide titled {AFTER!r}")
    kit = Kit(prs)

    s1 = kit.new_slide(FIRST)
    kit.bullets(s1, IS, 0.7, 1.6, 11.9, 4.2, size=20)
    kit.bullets(s1, ["Newton is for physics: it steps a simulated world forward."], 0.7, 5.9, 11.9, 0.5,
                size=20, name="takeaway", plain=True)
    kit.bullets(kit.new_slide("What Newton is not"), IS_NOT, 0.7, 1.7, 11.9, 4.6, size=20)
    s3 = kit.new_slide("Newton and its neighbours")
    kit.table(s3, NEIGHBOURS["cols"], NEIGHBOURS["rows"], 0.7, 1.7, row_h=0.85, size=18, name="neighbours")
    s4 = kit.new_slide("Start with the Newton examples")
    h = kit.command_box(s4, EXAMPLES_CMD)
    kit.bullets(s4, EXAMPLES, 0.7, 1.7 + h + 0.35, 11.9, 3.5, size=18)

    # Keep the agenda slide's total in step with the script's time budget (~32 min talk + ~7.5 min runs).
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.name == "stat talk plus simulation":
                sh._element.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}t")[0].text = "~40 min"

    kit.move_new_after(prs, titles.index(AFTER), 4)
    prs.save(args.out)
    print(f"saved {args.out}: {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
