"""Insert the solver-background slides into the demo deck.

    uvx --from python-pptx python slides/add_solver_slides.py in.pptx slides-<timestamp>-<message>.pptx

Adds four slides after "Two solvers, one scene": Newton's solver lineup (table), SolverMuJoCo,
SolverKamino, and solver coupling (proxy vs ADMM). Content comes from the Newton docs
(solvers/index, solvers/mujoco, solvers/kamino, concepts/coupling). The slides reuse the deck's
own title, bullet and table styling by cloning shapes from existing slides. Refuses to run on a
deck that already has them.
"""

import argparse
import copy
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Inches

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
AFTER = "Two solvers, one scene"
FIRST = "Newton's solvers"

TABLE = {
    "cols": [2.6, 2.6, 2.7, 4.0],
    "rows": [
        ("Solver", "Coordinates", "Handles", "Notes"),
        ("SolverMuJoCo", "generalized", "rigid, articulated", "MuJoCo Warp; our arm"),
        ("SolverFeatherstone", "generalized", "rigid, particles, soft", "basic differentiability"),
        ("SolverXPBD", "maximal", "rigid, particles, cloth", "implicit; soft bodies experimental"),
        ("SolverSemiImplicit", "maximal", "rigid, particles, soft", "basic differentiability"),
        ("SolverVBD", "maximal", "rigid, particles, cloth, soft", "implicit, experimental; our cable"),
        ("SolverKamino", "maximal", "rigid, articulated", "kinematic loops; experimental"),
        ("SolverStyle3D", "-", "particles, cloth", "no joints"),
        ("SolverImplicitMPM", "-", "particles", "no rigid bodies or joints"),
    ],
}

MUJOCO = [
    "Wraps MuJoCo Warp behind Newton's solver interface: GPU, many worlds",
    "Articulations in generalized coordinates: the best fit for a robot arm",
    "Defaults to the implicitfast integrator for stability with stiff PD gains",
    "use_mujoco_contacts=False feeds in Newton's collision pipeline instead (we do this)",
    "All worlds must have identical entity counts (why we pad the cable)",
    "No ROD joints, cloth, particles or soft bodies; non-convex meshes are convex-hulled",
]
KAMINO = [
    "Rigid mechanisms in maximal coordinates: kinematic loops, joint limits, hard frictional contacts",
    "Two backends: padmm (default, robust) and dvi (faster, less accurate with many active contacts)",
    "Per-DoF actuation, armature, damping and Coulomb friction",
    "Not for particles or deformables",
    "Experimental (BETA 1): the API and behavior may change without notice",
    "Not used in this demo; worth a look if closed chains matter",
]
PROXY = [
    "The source solver owns the real body; the destination sees a proxy with scaled inertia",
    "mode=\"lagged\": rewind the proxy by the last feedback, then step (our choice)",
    "mode=\"staggered\": use the source's current state directly",
    "mass_scale and proxy_relaxation set the stability",
    "Up to two entries; supports MPM transfer proxies",
]
ADMM = [
    "Symmetric: interface rows between bodies or particles owned by different solvers",
    "Cross-solver BALL, FIXED and REVOLUTE joints, body-particle attachments, contacts",
    "A fixed iteration count (graph-capture friendly); tune rho, gamma and baumgarte",
    "More parameters and fewer supported row types than proxy coupling",
]


def title_shape(slide):
    return next((s for s in slide.shapes if s.is_placeholder and s.placeholder_format.type == PP_PLACEHOLDER.TITLE), None)


def title_of(slide) -> str:
    sp = title_shape(slide)
    return sp.text_frame.text if sp is not None else ""


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck", type=Path)
    ap.add_argument("out", type=Path)
    args = ap.parse_args()
    if args.out.resolve() == args.deck.resolve():
        ap.error("out must differ from deck (keep the original)")

    prs = Presentation(args.deck)
    slides = list(prs.slides)
    titles = [title_of(s) for s in slides]
    if FIRST in titles:
        ap.error(f"deck already has the {FIRST!r} slide")
    if AFTER not in titles:
        ap.error(f"no slide titled {AFTER!r}")

    ref = slides[titles.index(AFTER)]
    layout = ref.slide_layout
    ref_title = title_shape(ref)
    ref_bullets = next(s for s in ref.shapes if s.name == "bullets")
    ref_num = next(s for s in ref.shapes if s.name.startswith("Slide Number"))
    ref_table = next(s for s in next(s for s in slides if title_of(s) == "Solver tuning").shapes if s.name == "tuning")
    tbl = ref_table._element.find(f".//{A}tbl")
    head_tc = copy.deepcopy(tbl.findall(f"{A}tr")[0].findall(f"{A}tc")[0])
    body_tc = copy.deepcopy(tbl.findall(f"{A}tr")[1].findall(f"{A}tc")[0])
    para_tpl = copy.deepcopy(ref_bullets._element.find(f".//{A}p"))

    def new_slide(title):
        s = prs.slides.add_slide(layout)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        for src, text in ((ref_title, title), (ref_num, None)):
            el = copy.deepcopy(src._element)
            if text is not None:
                el.findall(f".//{A}t")[0].text = text
            s.shapes._spTree.append(el)
        return s

    def bullets(slide, items, x, y, w, h, heading=None, size=18, name="bullets", plain=False):
        el = copy.deepcopy(ref_bullets._element)
        el.find(".//{*}cNvPr").set("name", name)
        el.find(".//{*}cNvPr").set("id", str(100 + len(slide.shapes)))
        slide.shapes._spTree.append(el)
        shape = slide.shapes[-1]
        shape.left, shape.top, shape.width, shape.height = Inches(x), Inches(y), Inches(w), Inches(h)
        body = el.find(f".//{{*}}txBody")
        for p in body.findall(f"{A}p"):
            body.remove(p)
        if heading:
            p = copy.deepcopy(para_tpl)
            p.find(f"{A}pPr").set("marL", "0")
            p.find(f"{A}pPr").set("indent", "0")
            p.find(f"{A}pPr").remove(p.find(f"{A}pPr").find(f"{A}buChar"))
            p.find(f"{A}pPr").append(p.makeelement(f"{A}buNone", {}))
            r = p.find(f"{A}r")
            r.find(f"{A}t").text = heading
            r.find(f"{A}rPr").set("b", "1")
            r.find(f"{A}rPr").set("sz", str((size + 4) * 100))
            body.append(p)
        for text in items:
            p = copy.deepcopy(para_tpl)
            r = p.find(f"{A}r")
            r.find(f"{A}t").text = text
            r.find(f"{A}rPr").set("sz", str(size * 100))
            if plain:  # a note, not a list item
                ppr = p.find(f"{A}pPr")
                ppr.set("marL", "0")
                ppr.set("indent", "0")
                ppr.remove(ppr.find(f"{A}buChar"))
                ppr.append(p.makeelement(f"{A}buNone", {}))
            body.append(p)

    def table(slide, spec, x, y, row_h=0.46, size=14):
        shape = slide.shapes.add_table(len(spec["rows"]), len(spec["cols"]), Inches(x), Inches(y),
                                       Inches(sum(spec["cols"])), Inches(row_h * len(spec["rows"])))
        shape.name = "solvers"
        gf = shape._element
        t = gf.find(f".//{A}tbl")
        for child in list(t):
            t.remove(child)
        grid = t.makeelement(f"{A}tblGrid", {})
        for w in spec["cols"]:
            grid.append(grid.makeelement(f"{A}gridCol", {"w": str(int(Inches(w)))}))
        t.append(t.makeelement(f"{A}tblPr", {}))
        t.append(grid)
        for i, row in enumerate(spec["rows"]):
            tr = t.makeelement(f"{A}tr", {"h": str(int(Inches(row_h)))})
            for text in row:
                tc = copy.deepcopy(head_tc if i == 0 else body_tc)
                tc.findall(f".//{A}t")[0].text = text
                for rpr in tc.findall(f".//{A}rPr"):
                    rpr.set("sz", str(size * 100))
                tr.append(tc)
            t.append(tr)

    s1 = new_slide(FIRST)
    table(s1, TABLE, 0.7, 1.5)
    bullets(s1, ["Pick the coordinate system first: generalized for robots, maximal for the rest. "
                 "Experimental: VBD and Kamino."], 0.7, 5.85, 11.9, 0.8, size=16, name="takeaway", plain=True)

    s2 = new_slide("SolverMuJoCo: the arm")
    bullets(s2, MUJOCO, 0.7, 1.7, 11.9, 4.6, size=20)
    s3 = new_slide("SolverKamino: for closed chains")
    bullets(s3, KAMINO, 0.7, 1.7, 11.9, 4.6, size=20)
    s4 = new_slide("Coupling: proxy vs ADMM")
    bullets(s4, PROXY, 0.7, 1.6, 5.8, 4.8, heading="Proxy (our demo)", size=17, name="proxy")
    bullets(s4, ADMM, 6.85, 1.6, 5.8, 4.8, heading="ADMM", size=17, name="admm")
    bullets(s4, ["Experimental API: it may change without notice."], 0.7, 6.35, 11.9, 0.4, size=14, name="note", plain=True)

    lst = prs.slides._sldIdLst
    ids = list(lst)[-4:]
    anchor = list(lst)[titles.index(AFTER)]
    for sid in ids:
        lst.remove(sid)
        anchor.addnext(sid)
        anchor = sid
    prs.save(args.out)
    print(f"saved {args.out}: {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
