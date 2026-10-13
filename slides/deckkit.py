"""Helpers shared by the slide-adding scripts: clone the deck's own title, bullet and table
shapes so new slides match the existing ones."""

import copy

from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Inches

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def title_shape(slide):
    """The title placeholder (the deck's titles use idx 100, so shapes.title misses them)."""
    return next((s for s in slide.shapes if s.is_placeholder and s.placeholder_format.type == PP_PLACEHOLDER.TITLE), None)


def title_of(slide) -> str:
    sp = title_shape(slide)
    return sp.text_frame.text if sp is not None else ""


class Kit:
    """Builds slides by cloning shapes from the deck's "Two solvers, one scene" (title, bullets,
    slide number) and "Solver tuning" (table) and "Domain randomization" (command box) slides."""

    def __init__(self, prs):
        self.prs = prs
        by_title = {title_of(s): s for s in prs.slides}
        ref = by_title["Two solvers, one scene"]
        self.layout = ref.slide_layout
        self.ref_title = title_shape(ref)
        self.ref_bullets = next(s for s in ref.shapes if s.name == "bullets")
        self.ref_num = next(s for s in ref.shapes if s.name.startswith("Slide Number"))
        tuning = next(s for s in by_title["Solver tuning"].shapes if s.name == "tuning")
        tbl = tuning._element.find(f".//{A}tbl")
        self.head_tc = copy.deepcopy(tbl.findall(f"{A}tr")[0].findall(f"{A}tc")[0])
        self.body_tc = copy.deepcopy(tbl.findall(f"{A}tr")[1].findall(f"{A}tc")[0])
        self.para_tpl = copy.deepcopy(self.ref_bullets._element.find(f".//{A}p"))
        rand = by_title["Domain randomization"]
        self.ref_cmd_bg = next(s for s in rand.shapes if s.name == "cmd bg")
        self.ref_cmd = next(s for s in rand.shapes if s.name == "cmd")

    def new_slide(self, title):
        s = self.prs.slides.add_slide(self.layout)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        for src, text in ((self.ref_title, title), (self.ref_num, None)):
            el = copy.deepcopy(src._element)
            if text is not None:
                el.findall(f".//{A}t")[0].text = text
            s.shapes._spTree.append(el)
        return s

    def bullets(self, slide, items, x, y, w, h, heading=None, size=18, name="bullets", plain=False):
        """A text box of bullet items (``plain``: no bullet markers, for a note)."""
        el = copy.deepcopy(self.ref_bullets._element)
        el.find(".//{*}cNvPr").set("name", name)
        el.find(".//{*}cNvPr").set("id", str(100 + len(slide.shapes)))
        slide.shapes._spTree.append(el)
        shape = slide.shapes[-1]
        shape.left, shape.top, shape.width, shape.height = Inches(x), Inches(y), Inches(w), Inches(h)
        body = el.find(".//{*}txBody")
        for p in body.findall(f"{A}p"):
            body.remove(p)

        def no_bullet(p):
            ppr = p.find(f"{A}pPr")
            ppr.set("marL", "0")
            ppr.set("indent", "0")
            ppr.remove(ppr.find(f"{A}buChar"))
            ppr.append(p.makeelement(f"{A}buNone", {}))

        if heading:
            p = copy.deepcopy(self.para_tpl)
            no_bullet(p)
            r = p.find(f"{A}r")
            r.find(f"{A}t").text = heading
            r.find(f"{A}rPr").set("b", "1")
            r.find(f"{A}rPr").set("sz", str((size + 4) * 100))
            body.append(p)
        for text in items:
            p = copy.deepcopy(self.para_tpl)
            r = p.find(f"{A}r")
            r.find(f"{A}t").text = text
            r.find(f"{A}rPr").set("sz", str(size * 100))
            if plain:
                no_bullet(p)
            body.append(p)

    def table(self, slide, cols, rows, x, y, row_h=0.46, size=14, name="table"):
        """A table in the deck's style: ``cols`` are widths in inches, ``rows[0]`` the header."""
        shape = slide.shapes.add_table(len(rows), len(cols), Inches(x), Inches(y),
                                       Inches(sum(cols)), Inches(row_h * len(rows)))
        shape.name = name
        t = shape._element.find(f".//{A}tbl")
        for child in list(t):
            t.remove(child)
        grid = t.makeelement(f"{A}tblGrid", {})
        for w in cols:
            grid.append(grid.makeelement(f"{A}gridCol", {"w": str(int(Inches(w)))}))
        t.append(t.makeelement(f"{A}tblPr", {}))
        t.append(grid)
        for i, row in enumerate(rows):
            tr = t.makeelement(f"{A}tr", {"h": str(int(Inches(row_h)))})
            for text in row:
                tc = copy.deepcopy(self.head_tc if i == 0 else self.body_tc)
                tc.findall(f".//{A}t")[0].text = text
                for rpr in tc.findall(f".//{A}rPr"):
                    rpr.set("sz", str(size * 100))
                tr.append(tc)
            t.append(tr)

    def command_box(self, slide, lines, x=0.7, y=1.7, w=11.93):
        """A command box in the deck's style (cloned from "Domain randomization"): gray ``$``,
        blue ``--flags``, orange numbers, gray ``#`` comments."""
        bg, cmd = self.ref_cmd_bg, self.ref_cmd
        for src in (bg, cmd):
            slide.shapes._spTree.append(copy.deepcopy(src._element))
        bg_s, cmd_s = slide.shapes[-2], slide.shapes[-1]
        h = 0.3 * len(lines) + 0.3
        bg_s.left, bg_s.top, bg_s.width, bg_s.height = Inches(x), Inches(y), Inches(w), Inches(h)
        cmd_s.left, cmd_s.top, cmd_s.width, cmd_s.height = Inches(x + 0.25), Inches(y + 0.12), Inches(w - 0.5), Inches(h - 0.24)
        paras = cmd_s._element.findall(f".//{A}p")
        runs = paras[1].findall(f"{A}r")  # templates: gray "$ ", dark text, blue flag, orange value
        gray, dark, blue, orange = runs[0], runs[1], runs[3], runs[-1]
        tpl = copy.deepcopy(paras[1])
        body = paras[0].getparent()
        for p in paras:
            body.remove(p)
        for line in lines:
            p = copy.deepcopy(tpl)
            for r in p.findall(f"{A}r"):
                p.remove(r)
            end = p.find(f"{A}endParaRPr")
            code, _, comment = line.partition("   #")
            pieces = [(gray, "$ ")]
            for tok in code[2:].split(" "):
                kind = blue if tok.startswith("--") else orange if tok.replace(".", "").isdigit() else dark
                if len(pieces) > 1:
                    pieces.append((dark, " "))
                pieces.append((kind, tok))
            if comment:
                pieces.append((gray, "   #" + comment))
            for tpl_run, text in pieces:
                r = copy.deepcopy(tpl_run)
                r.findall(f"{A}t")[0].text = text
                end.addprevious(r) if end is not None else p.append(r)
            body.append(p)
        return h

    @staticmethod
    def move_new_after(prs, anchor_index, count):
        """Move the last ``count`` slides (just added) to follow the slide at ``anchor_index``
        (an index into the slide list as it was before adding them)."""
        lst = prs.slides._sldIdLst
        anchor = list(lst)[anchor_index]
        for sid in list(lst)[-count:]:
            lst.remove(sid)
            anchor.addnext(sid)
            anchor = sid
