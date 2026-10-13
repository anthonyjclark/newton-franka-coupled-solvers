"""Insert simulation-video slides into the demo deck.

    uvx --from python-pptx python slides/add_videos.py in.pptx slides-<timestamp>-<message>.pptx [--videos slides/videos]

Videos come from slides/make_videos.sh. Each new slide goes right after the slide whose
title is listed below, reusing the "Recording:" slide layout. Poster frames are extracted
with ffmpeg. Entries whose slide already exists are skipped, so it can top up a deck that
already has some "Video:" slides; it fails if there is nothing to add.
"""

import argparse
import copy
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Inches

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

# (title of the slide to insert after, new title, caption, [(video name, label)])
NEW = [
    ("First demo: a single arm", "Video: a single arm",
     "Same run as the live demo: one world, 7 s task rendered at 60 fps.",
     [("single-arm", "")]),
    ("Recording: 16 worlds", "Video: 16 worlds",
     "16 independent Franka + cable worlds, simulated in parallel on one GPU.",
     [("multi-world", "")]),
    ("Change the task", "Video: speed and place target",
     "Left: --sim.speed 2.0. Right: --sim.place-x 0.3 --sim.place-y -0.2.",
     [("fast", "2x speed"), ("place-target", "place target (0.3, -0.2)")]),
    ("Change the cable", "Video: cable variants",
     "Same task, three cables: 31 segments / 0.45 m, stiff bend, soft bend.",
     [("long-cable", "longer cable"), ("stiff-cable", "bend stiffness 5e-2"),
      ("soft-cable", "bend stiffness 5e-5")]),
    ("Recording: randomization", "Video: randomization",
     "8 worlds with --sim.rand.all: each world gets its own friction, mass, gains and timing.",
     [("randomized", "")]),
    ("Video: randomization", "Video: expensive randomization",
     "8 worlds with --sim.rand.all --sim.rand.expensive: every world also gets its own cable "
     "length, radius, segment count, stiffness and yaw.",
     [("randomized-expensive", "")]),
]

X0, W_ALL, Y0, H_ALL, GAP = 0.7, 11.9, 1.6, 4.3, 0.3  # content area of the recording slides


def poster(video: Path, dest: Path) -> Path:
    """Extract a frame about 2/3 of the way through the video."""
    dur = float(subprocess.run(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(video)],
        capture_output=True, text=True, check=True).stdout)
    png = dest / f"{video.stem}.png"
    subprocess.run(["ffmpeg", "-loglevel", "error", "-y", "-ss", str(dur * 0.64), "-i", str(video),
                    "-frames:v", "1", str(png)], check=True)
    return png


def title_shape(slide):
    """The title placeholder (the deck's titles use idx 100, so shapes.title misses them)."""
    return next((s for s in slide.shapes if s.is_placeholder and s.placeholder_format.type == PP_PLACEHOLDER.TITLE), None)


def title_of(slide) -> str:
    sp = title_shape(slide)
    return sp.text_frame.text if sp is not None else ""


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck", type=Path)
    ap.add_argument("out", type=Path)
    ap.add_argument("--videos", type=Path, default=Path("slides/videos"))
    args = ap.parse_args()
    if args.out.resolve() == args.deck.resolve():
        ap.error("out must differ from deck (keep the original)")

    prs = Presentation(args.deck)
    slides = list(prs.slides)
    titles = [title_of(s) for s in slides]
    todo = [e for e in NEW if e[1] not in titles]
    if not todo:
        ap.error("deck already has every Video: slide")
    known = set(titles) | {e[1] for e in todo}
    missing = [a for a, *_ in todo if a not in known]
    if missing:
        ap.error(f"no slide titled {missing}")
    for *_, vids in todo:
        for name, _ in vids:
            if not (args.videos / f"{name}.mp4").is_file():
                ap.error(f"missing {args.videos / (name + '.mp4')} (run slides/make_videos.sh)")

    ref = next(s for s in slides if title_of(s).startswith("Recording:"))
    layout = ref.slide_layout
    ref_title = title_shape(ref)
    ref_caption = next(s for s in ref.shapes if s.name == "caption")
    ref_num = next(s for s in ref.shapes if s.name.startswith("Slide Number"))
    muted = copy.deepcopy(ref_caption._element.find(f".//{A}solidFill"))

    tmp = Path(tempfile.mkdtemp(prefix="posters-"))

    def make_slide(title, caption, vids):
        s = prs.slides.add_slide(layout)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        for src, text in ((ref_title, title), (ref_caption, caption), (ref_num, None)):
            el = copy.deepcopy(src._element)
            if text is not None:
                el.findall(f".//{A}t")[0].text = text
            s.shapes._spTree.append(el)
        labelled = any(lbl for _, lbl in vids)
        pad = 0.4 if labelled else 0
        n = len(vids)
        w = (W_ALL - GAP * (n - 1)) / n
        h = min(w * 9 / 16, H_ALL - pad)
        w = h * 16 / 9
        x = X0 + (W_ALL - (w * n + GAP * (n - 1))) / 2
        y = Y0 + (H_ALL - h - pad) / 2
        for name, label in vids:
            video = args.videos / f"{name}.mp4"
            m = s.shapes.add_movie(str(video), Inches(x), Inches(y), Inches(w), Inches(h),
                                   poster_frame_image=str(poster(video, tmp)), mime_type="video/mp4")
            m.name = f"video {name}"
            m._element.nvPicPr.cNvPr.set("descr", f"{label or title} simulation video")
            if label:
                tb = s.shapes.add_textbox(Inches(x), Inches(y + h + 0.05), Inches(w), Inches(0.3))
                tb.name = f"label {name}"
                tf = tb.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.text = label
                run = tf.paragraphs[0].runs[0]
                run.font.size = Inches(12 / 72)
                run._r.get_or_add_rPr().append(copy.deepcopy(muted))
            x += w + GAP

    lst = prs.slides._sldIdLst
    anchors = dict(zip(titles, list(lst)))  # slide ids by title (new slides are added as they go)
    for after, title, caption, vids in todo:
        make_slide(title, caption, vids)
        new_id = list(lst)[-1]
        lst.remove(new_id)
        anchors[after].addnext(new_id)
        anchors[title] = new_id
    prs.save(args.out)
    print(f"saved {args.out}: {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
